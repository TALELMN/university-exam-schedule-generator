from openpyxl import load_workbook

from exam_scheduler.config import AppConfig
from exam_scheduler.output.generate import generate_outputs
from exam_scheduler.output.common import safe_filename, safe_sheet_name
from tests.test_matching import _sample_result


def test_generates_excel_pdf_html_and_validation_report(tmp_path):
    result = _sample_result()
    files = generate_outputs(result, tmp_path, AppConfig())
    names = {path.name for path in files}
    assert {"University_Exam_Schedules.xlsx", "ALL_LEVELS.xlsx", "University_Exam_Schedules.pdf",
            "University_Exam_Schedules.html", "Validation_Report.txt", "Freshman_5.xlsx"} <= names
    workbook = load_workbook(tmp_path / "University_Exam_Schedules.xlsx", read_only=True)
    sheet = workbook["Freshman 5"]
    codes = [sheet.cell(row, 4).value for row in range(2, sheet.max_row + 1)]
    assert codes == ["ECO171", "ENG111", "ECE143", "CS101", "PHYS151", "MATH141"]
    assert (tmp_path / "Validation_Report.txt").read_text(encoding="utf-8").find("CS999") >= 0


def test_level_names_are_safe_for_filenames_and_excel_sheets():
    assert safe_filename("Junior/CSE: A") == "Junior_CSE_A"
    assert safe_sheet_name("Junior/CSE: A") == "Junior_CSE_ A"
