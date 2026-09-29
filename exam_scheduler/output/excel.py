"""Styled Excel workbooks for per-level and university-wide schedules."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

from exam_scheduler.matching.validator import validation_counts
from exam_scheduler.models import AnalysisResult
from exam_scheduler.output.common import export_directory, rows_for_level, safe_filename, safe_sheet_name, time_label
from exam_scheduler.utils.dates import display_date

LEVEL_COLUMNS = ["Date", "Day", "Time", "Course Code", "Course Name", "Professor(s)", "Status"]


def _level_rows(result: AnalysisResult, level: str) -> list[dict[str, str]]:
    output = []
    for row in rows_for_level(result, level):
        output.append({
            "Date": display_date(row.date), "Day": row.day or "—",
            "Time": time_label(row.start_time, row.end_time), "Course Code": row.code,
            "Course Name": row.course_name or ("Unknown" if row.status == "Missing" else ""),
            "Professor(s)": row.professors or "", "Status": row.status,
        })
    return output


def _master_rows(result: AnalysisResult) -> list[dict[str, str]]:
    levels_by_code: dict[str, list[str]] = {}
    for row in result.matches:
        levels_by_code.setdefault(row.code, [])
        if row.level not in levels_by_code[row.code]:
            levels_by_code[row.code].append(row.level)
    rows = []
    for code, exam in result.exams.items():
        rows.append({
            "_sort_date": exam.date or "9999-99-99", "Date": display_date(exam.date), "Day": exam.day or "—",
            "Time": time_label(exam.start_time, exam.end_time), "Course Code": code,
            "Course Name": exam.course_name or "", "Level(s)": ", ".join(sorted(levels_by_code.get(code, []))),
        })
    rows.sort(key=lambda row: (row["_sort_date"], row["Time"], row["Course Code"]))
    return [{key: value for key, value in row.items() if key != "_sort_date"} for row in rows]


def _style_workbook(path: Path) -> None:
    from openpyxl import load_workbook
    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="17365D")
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2" if worksheet.max_row else "A1"
        worksheet.sheet_view.showGridLines = False
        if worksheet.max_row >= 1 and worksheet.max_column >= 1:
            worksheet.auto_filter.ref = worksheet.dimensions
            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(vertical="center")
        for column in worksheet.columns:
            values = [len(str(cell.value or "")) for cell in column]
            width = min(max(max(values, default=10) + 2, 12), 52)
            worksheet.column_dimensions[column[0].column_letter].width = width
    workbook.save(path)


def generate_schedule_workbook(result: AnalysisResult, output_dir: str | Path) -> list[Path]:
    """Generate the consolidated workbook plus a standalone workbook per level."""
    directory = export_directory(output_dir)
    created: list[Path] = []
    combined = directory / "University_Exam_Schedules.xlsx"
    counts = validation_counts(result)
    summary_rows = [{"Metric": label, "Count": value} for label, value in counts.items()]
    summary_rows += [{"Metric": "Exam PDF", "Count": result.exam_source}]
    summary = pd.DataFrame(summary_rows)
    used_sheets: set[str] = {"summary"}
    with pd.ExcelWriter(combined, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Summary", index=False)
        for level in result.levels:
            base = safe_sheet_name(level.name)
            sheet = base
            suffix = 2
            while sheet.casefold() in used_sheets:
                tail = f"_{suffix}"
                sheet = base[:31 - len(tail)] + tail
                suffix += 1
            used_sheets.add(sheet.casefold())
            pd.DataFrame(_level_rows(result, level.name), columns=LEVEL_COLUMNS).to_excel(
                writer, sheet_name=sheet, index=False
            )
    _style_workbook(combined)
    created.append(combined)

    master_path = directory / "ALL_LEVELS.xlsx"
    with pd.ExcelWriter(master_path, engine="openpyxl") as writer:
        pd.DataFrame(_master_rows(result), columns=["Date", "Day", "Time", "Course Code", "Course Name", "Level(s)"]).to_excel(
            writer, sheet_name="All Exams", index=False
        )
        pd.DataFrame(summary_rows).to_excel(writer, sheet_name="Summary", index=False)
    _style_workbook(master_path)
    created.append(master_path)

    used_files = {combined.name.casefold(), master_path.name.casefold()}
    for level in result.levels:
        base = safe_filename(level.name)
        filename, suffix = f"{base}.xlsx", 2
        while filename.casefold() in used_files:
            filename = f"{base}_{suffix}.xlsx"
            suffix += 1
        used_files.add(filename.casefold())
        path = directory / filename
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            pd.DataFrame(_level_rows(result, level.name), columns=LEVEL_COLUMNS).to_excel(
                writer, sheet_name="Exam Schedule", index=False
            )
        _style_workbook(path)
        created.append(path)
    return created
